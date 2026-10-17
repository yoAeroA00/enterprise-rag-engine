import io
import logging
from pathlib import Path
from typing import Any

import PyPDF2
from docx import Document
from langchain.schema import Document as LangchainDocument
from langchain.text_splitter import RecursiveCharacterTextSplitter

from config import Config

logger = logging.getLogger(__name__)


class DocumentProcessor:
    """Handle document processing and text extraction"""

    def __init__(self):
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=Config.CHUNK_SIZE,
            chunk_overlap=Config.CHUNK_OVERLAP,
            length_function=len,
            separators=["\n\n", "\n", " ", ""],
        )

    def validate_file(self, file) -> dict[str, Any]:
        """Validate uploaded file"""
        validation_result = {"valid": True, "error": None, "warnings": []}

        try:
            # Check file size
            if hasattr(file, "size") and file.size > Config.MAX_FILE_SIZE:
                validation_result["valid"] = False
                validation_result["error"] = (
                    f"File size ({file.size / (1024 * 1024):.1f}MB) exceeds maximum allowed size ({Config.MAX_FILE_SIZE / (1024 * 1024):.1f}MB)"
                )
                return validation_result

            # Check file extension
            file_extension = Path(file.name).suffix.lower()
            if file_extension not in Config.ALLOWED_EXTENSIONS:
                validation_result["valid"] = False
                validation_result["error"] = (
                    f"File type '{file_extension}' not supported. Allowed types: {', '.join(Config.ALLOWED_EXTENSIONS)}"
                )
                return validation_result

            # Check if file is empty
            if hasattr(file, "size") and file.size == 0:
                validation_result["valid"] = False
                validation_result["error"] = "File is empty"
                return validation_result

        except Exception as e:
            validation_result["valid"] = False
            validation_result["error"] = f"File validation error: {e}"
            logger.error(f"File validation error: {e}")

        return validation_result

    def extract_text_from_pdf(self, file) -> str:
        """Extract text from PDF file"""
        try:
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(file.read()))
            text = ""

            for page_num, page in enumerate(pdf_reader.pages):
                try:
                    page_text = page.extract_text()
                    if page_text.strip():  # Only add non-empty pages
                        text += f"\n--- Page {page_num + 1} ---\n"
                        text += page_text
                except Exception as e:
                    logger.warning(
                        f"Error extracting text from page {page_num + 1}: {e}",
                        exc_info=True,
                    )
                    continue

            if not text.strip():
                raise ValueError("No text could be extracted from the PDF")

            return text.strip()

        except Exception as e:
            logger.error(f"Error extracting text from PDF: {e}")
            raise ValueError(f"Failed to extract text from PDF: {e}") from e

    def extract_text_from_docx(self, file) -> str:
        """Extract text from DOCX file"""
        try:
            doc = Document(io.BytesIO(file.read()))
            text = ""

            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text += paragraph.text + "\n"

            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        text += " | ".join(row_text) + "\n"

            if not text.strip():
                raise ValueError("No text could be extracted from the DOCX file")

            return text.strip()

        except Exception as e:
            logger.error(f"Error extracting text from DOCX: {e}")
            raise ValueError(f"Failed to extract text from DOCX: {e}") from e

    def extract_text_from_txt(self, file) -> str:
        """Extract text from TXT file"""
        try:
            # Try different encodings
            encodings = ["utf-8", "utf-16", "latin-1", "cp1252"]

            for encoding in encodings:
                try:
                    file.seek(0)  # Reset file pointer
                    content = file.read()
                    text = (
                        content.decode(encoding)
                        if isinstance(content, bytes)
                        else content
                    )

                    if text.strip():
                        return text.strip()
                except (UnicodeDecodeError, UnicodeError):
                    continue

            raise ValueError(
                "Could not decode the text file with any supported encoding"
            )

        except Exception as e:
            logger.error(f"Error extracting text from TXT: {e}")
            raise ValueError(f"Failed to extract text from TXT file: {e}") from e

    def extract_text(self, file) -> str:
        """Extract text from uploaded file based on file type"""
        file_extension = Path(file.name).suffix.lower()

        try:
            if file_extension == ".pdf":
                return self.extract_text_from_pdf(file)
            elif file_extension == ".docx":
                return self.extract_text_from_docx(file)
            elif file_extension in [".txt", ".md"]:
                return self.extract_text_from_txt(file)
            else:
                raise ValueError(f"Unsupported file type: {file_extension}")

        except Exception as e:
            logger.error(f"Error extracting text from {file.name}: {e}")
            raise

    def split_text(
        self, text: str, metadata: dict | None = None
    ) -> list[LangchainDocument]:
        """Split text into chunks"""
        try:
            if not text or not text.strip():
                raise ValueError("Text is empty or contains only whitespace")

            # Create documents with metadata
            documents = [LangchainDocument(page_content=text, metadata=metadata or {})]

            # Split the documents
            split_docs = self.text_splitter.split_documents(documents)

            if not split_docs:
                raise ValueError("Text splitting resulted in no documents")

            # Add chunk information to metadata
            for i, doc in enumerate(split_docs):
                doc.metadata.update(
                    {
                        "chunk_id": i,
                        "total_chunks": len(split_docs),
                        "chunk_size": len(doc.page_content),
                    }
                )

            logger.info(f"Split text into {len(split_docs)} chunks")
            return split_docs

        except Exception as e:
            logger.error(f"Error splitting text: {e}")
            raise ValueError(f"Failed to split text: {e}") from e

    def process_file(
        self, file, additional_metadata: dict | None = None
    ) -> list[LangchainDocument]:
        """Process uploaded file and return document chunks"""
        try:
            # Validate file
            validation = self.validate_file(file)
            if not validation["valid"]:
                raise ValueError(validation["error"])

            # Extract text
            text = self.extract_text(file)

            # Prepare metadata
            metadata = {
                "filename": file.name,
                "file_type": Path(file.name).suffix.lower(),
                "file_size": getattr(file, "size", 0),
            }

            if additional_metadata:
                metadata.update(additional_metadata)

            # Split text into chunks
            documents = self.split_text(text, metadata)

            logger.info(
                f"Successfully processed file '{file.name}' into {len(documents)} chunks"
            )
            return documents

        except Exception as e:
            logger.error(f"Error processing file '{file.name}': {e}")
            raise
